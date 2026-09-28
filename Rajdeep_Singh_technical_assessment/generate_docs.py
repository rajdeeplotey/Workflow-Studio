# generate_docs.py - Generates PDF and DOCX user guide for VectorShift Workflow Studio

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def create_docx(filename):
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("VectorShift Workflow Studio")
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Comprehensive Features & User Documentation Guide")
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Sections Content Data
    sections_data = [
        ("1. Multi-Workflow Sheet Tabs (+ New Workflows & Rename)", [
            ("Create New Workflows (+ Button)", "Click the '+' icon on the bottom tab bar to create new workflow tabs with sequential auto-naming (e.g., Untitled 1, Untitled 2)."),
            ("Tab Switching & Navigation", "Click any workflow tab at the bottom to switch instantly between active workflows while preserving full canvas state."),
            ("Inline Tab Renaming", "Double-click any tab name to rename it inline. Press Enter or click outside to save the new name.")
        ]),
        ("2. Excel-Style Save or Don't Save Modal Prompt", [
            ("Unsaved Tab Close Handler", "When clicking the '✕' close button on an unsaved tab, an Excel-style confirmation modal appears asking to Save or Don't Save."),
            ("Save Action", "Triggers native laptop file save picker (.json) to save the workflow directly to your computer/desktop, marks the tab as saved, and closes cleanly."),
            ("Don't Save Action", "Discards unsaved changes and closes the tab immediately without popups or moving to trash."),
            ("Saved Tab 1-Click Close", "Saved tabs close cleanly with zero popups and zero clutter.")
        ]),
        ("3. Open / Import Existing JSON Files (📂 Icon)", [
            ("Import Existing Files", "Click the folder icon (📂) on the bottom tab bar to select and parse any JSON workflow file from your computer/laptop directly into active studio sheets.")
        ]),
        ("4. Hand Pan Mode vs. Arrow Selection Mode", [
            ("Hand Pan Mode (🖐️)", "Click and drag anywhere on the canvas background to pan smoothly. Connect node handles by dragging lines between dots. Double-click any edge connection line to delete it."),
            ("Arrow Selection Mode (↖️)", "Click and drag a box over multiple nodes to batch select, move, or copy. Double-click edge deletion is suppressed in Arrow mode to prevent accidental deletion during box selection.")
        ]),
        ("5. Dual-Layer Laptop File Deletion & 30-Day Studio Trash", [
            ("Laptop File Deletion Auto-Detection", "If a saved JSON file is deleted from your computer/desktop or recycle bin, VectorShift Studio automatically detects the missing file when the window regains focus."),
            ("30-Day Studio Trash Retention", "Missing or deleted workflows are safely preserved in Studio Trash for 30 days before auto-purging."),
            ("Trash Management Submenu", "Click Three-Dots Menu ➔ Trash to view deleted items, restore them back to active tabs, or permanently empty trash.")
        ]),
        ("6. Copy, Paste & Undo/Redo Engine", [
            ("Copy & Paste (Ctrl+C / Ctrl+V)", "Select nodes and press Ctrl+C (Cmd+C) to copy. Press Ctrl+V (Cmd+V) to paste copied nodes onto the canvas with offset positioning."),
            ("Undo & Redo (Ctrl+Z / Ctrl+Y)", "Full 50-step state history. Press Ctrl+Z to undo and Ctrl+Y (or Ctrl+Shift+Z) to redo any action.")
        ]),
        ("7. Deleting Nodes & Edges (Cross Button / Double Click)", [
            ("Node Removal (Cross Button ✕)", "Click the '✕' button on the top-right corner of any node card to delete it instantly."),
            ("Edge Removal (Double Click)", "Double-click any connection line on the canvas to delete it instantly.")
        ]),
        ("8. Appearance & Theme Engine (Default, Light, Dark)", [
            ("Theme Submenu", "Click Three-Dots Menu ➔ Appearance to toggle between Default (VectorShift Official Brand), Light Mode, and Dark Mode with smooth CSS transitions.")
        ]),
        ("9. Touchpad & Touch Zoom Engine", [
            ("2-Finger Touchpad Zoom", "Slide 2 fingers up/down on your laptop touchpad to zoom in and out smoothly without needing Ctrl."),
            ("Touchpad Pinch Zoom", "Pinch in/out with 2 fingers on trackpad to zoom in/out."),
            ("Ctrl + Scroll Wheel", "Holding Ctrl + mouse wheel scroll continues to work as expected.")
        ]),
        ("10. Real Pipeline Execution Engine (Bonus Feature)", [
            ("Analyze Pipeline (/pipelines/parse)", "Calculates node count, edge count, and validates DAG cycle status."),
            ("Run Execution Engine (/pipelines/execute)", "Executes nodes topologically: substitutes {{variables}}, evaluates Math/Filter logic, awaits Timer delays, calls real/mock LLM API, and displays step-by-step execution logs and terminal outputs.")
        ])
    ]

    for title, items in sections_data:
        h = doc.add_heading(title, level=2)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

        for item_title, item_desc in items:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_after = Pt(4)
            r_b = p.add_run(f"• {item_title}: ")
            r_b.bold = True
            r_b.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            p.add_run(item_desc)

    doc.save(filename)
    print(f"DOCX saved to {filename}")

def create_pdf(filename):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1E293B'),
        alignment=1,
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#2563EB'),
        alignment=1,
        spaceAfter=18
    )

    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=12,
        spaceAfter=6,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        'DocBullet',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13.5,
        textColor=colors.HexColor('#334155'),
        leftIndent=12,
        spaceAfter=4
    )

    story = []

    story.append(Paragraph("VectorShift Workflow Studio", title_style))
    story.append(Paragraph("Comprehensive Features & User Documentation Guide", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E2E8F0'), spaceAfter=12))

    sections_data = [
        ("1. Multi-Workflow Sheet Tabs (+ New Workflows & Rename)", [
            ("Create New Workflows (+ Button)", "Click the '+' icon on the bottom tab bar to create new workflow tabs with sequential auto-naming (e.g., Untitled 1, Untitled 2)."),
            ("Tab Switching & Navigation", "Click any workflow tab at the bottom to switch instantly between active workflows while preserving full canvas state."),
            ("Inline Tab Renaming", "Double-click any tab name to rename it inline. Press Enter or click outside to save the new name.")
        ]),
        ("2. Excel-Style Save or Don't Save Modal Prompt", [
            ("Unsaved Tab Close Handler", "When clicking the '✕' close button on an unsaved tab, an Excel-style confirmation modal appears asking to Save or Don't Save."),
            ("Save Action", "Triggers native laptop file save picker (.json) to save the workflow directly to your computer/desktop, marks the tab as saved, and closes cleanly."),
            ("Don't Save Action", "Discards unsaved changes and closes the tab immediately without popups or moving to trash."),
            ("Saved Tab 1-Click Close", "Saved tabs close cleanly with zero popups and zero clutter.")
        ]),
        ("3. Open / Import Existing JSON Files (📂 Icon)", [
            ("Import Existing Files", "Click the folder icon (📂) on the bottom tab bar to select and parse any JSON workflow file from your computer/laptop directly into active studio sheets.")
        ]),
        ("4. Hand Pan Mode vs. Arrow Selection Mode", [
            ("Hand Pan Mode (🖐️)", "Click and drag anywhere on the canvas background to pan smoothly. Connect node handles by dragging lines between dots. Double-click any edge connection line to delete it."),
            ("Arrow Selection Mode (↖️)", "Click and drag a box over multiple nodes to batch select, move, or copy. Double-click edge deletion is suppressed in Arrow mode to prevent accidental deletion during box selection.")
        ]),
        ("5. Dual-Layer Laptop File Deletion & 30-Day Studio Trash", [
            ("Laptop File Deletion Auto-Detection", "If a saved JSON file is deleted from your computer/desktop or recycle bin, VectorShift Studio automatically detects the missing file when the window regains focus."),
            ("30-Day Studio Trash Retention", "Missing or deleted workflows are safely preserved in Studio Trash for 30 days before auto-purging."),
            ("Trash Management Submenu", "Click Three-Dots Menu ➔ Trash to view deleted items, restore them back to active tabs, or permanently empty trash.")
        ]),
        ("6. Copy, Paste & Undo/Redo Engine", [
            ("Copy & Paste (Ctrl+C / Ctrl+V)", "Select nodes and press Ctrl+C (Cmd+C) to copy. Press Ctrl+V (Cmd+V) to paste copied nodes onto the canvas with offset positioning."),
            ("Undo & Redo (Ctrl+Z / Ctrl+Y)", "Full 50-step state history. Press Ctrl+Z to undo and Ctrl+Y (or Ctrl+Shift+Z) to redo any action.")
        ]),
        ("7. Deleting Nodes & Edges (Cross Button / Double Click)", [
            ("Node Removal (Cross Button ✕)", "Click the '✕' button on the top-right corner of any node card to delete it instantly."),
            ("Edge Removal (Double Click)", "Double-click any connection line on the canvas to delete it instantly.")
        ]),
        ("8. Appearance & Theme Engine (Default, Light, Dark)", [
            ("Theme Submenu", "Click Three-Dots Menu ➔ Appearance to toggle between Default (VectorShift Official Brand), Light Mode, and Dark Mode with smooth CSS transitions.")
        ]),
        ("9. Touchpad & Touch Zoom Engine", [
            ("2-Finger Touchpad Zoom", "Slide 2 fingers up/down on your laptop touchpad to zoom in and out smoothly without needing Ctrl."),
            ("Touchpad Pinch Zoom", "Pinch in/out with 2 fingers on trackpad to zoom in/out."),
            ("Ctrl + Scroll Wheel", "Holding Ctrl + mouse wheel scroll continues to work as expected.")
        ]),
        ("10. Real Pipeline Execution Engine (Bonus Feature)", [
            ("Analyze Pipeline (/pipelines/parse)", "Calculates node count, edge count, and validates DAG cycle status."),
            ("Run Execution Engine (/pipelines/execute)", "Executes nodes topologically: substitutes {{variables}}, evaluates Math/Filter logic, awaits Timer delays, calls real/mock LLM API, and displays step-by-step execution logs and terminal outputs.")
        ])
    ]

    for title, items in sections_data:
        story.append(Paragraph(title, h2_style))
        for item_title, item_desc in items:
            text = f"<b><font color='#2563EB'>• {item_title}:</font></b> {item_desc}"
            story.append(Paragraph(text, bullet_style))
        story.append(Spacer(1, 4))

    doc.build(story)
    print(f"PDF saved to {filename}")

if __name__ == '__main__':
    root_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(root_dir, "VectorShift_Studio_Features_Guide.docx")
    pdf_path = os.path.join(root_dir, "VectorShift_Studio_Features_Guide.pdf")
    
    create_docx(docx_path)
    create_pdf(pdf_path)
